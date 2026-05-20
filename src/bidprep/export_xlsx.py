from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bidprep.schemas import MergeMeta, TenderExtraction

HEADER_FILL = PatternFill(start_color="66CCFF", end_color="66CCFF", fill_type="solid")
HEADER_FONT = Font(bold=True, color="000000")


def _style_header(ws, row: int = 1, cols: int = 5) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_materials_xlsx(extraction: TenderExtraction, out_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "材料清单"
    headers = ["序号", "材料类别", "材料名称", "是否具备", "备注"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, 5)
    for idx, m in enumerate(extraction.materials, start=1):
        r = idx + 1
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=m.category)
        ws.cell(row=r, column=3, value=m.name)
        ws.cell(row=r, column=4, value=m.status)
        remark = m.remark
        if m.source_page:
            remark = (remark + f"（约第{m.source_page}页）").strip("（）") if remark else f"约第{m.source_page}页"
        ws.cell(row=r, column=5, value=remark)
    if not extraction.materials:
        ws.cell(row=2, column=1, value=1)
        ws.cell(row=2, column=2, value="—")
        ws.cell(row=2, column=3, value="请根据招标文件人工补充材料行")
        ws.cell(row=2, column=4, value="☐ 需准备")
        ws.cell(row=2, column=5, value="可重新运行抽取或对照 PDF")

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col != 3 else 36

    ws2 = wb.create_sheet("盘点摘要")
    ws2.cell(row=1, column=1, value="已具备材料清单（客户填写文件名等）")
    ws2.cell(row=2, column=1, value="（在此下方逐条列出）")
    ws2.cell(row=4, column=1, value="仍需准备材料")
    ws2.cell(row=5, column=1, value="（在此下方逐条列出）")

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def export_fill_items_xlsx(extraction: TenderExtraction, meta: MergeMeta | None, out_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "填空项"
    headers = ["序号", "章节/段落", "字段名称", "说明", "必填", "招标文件摘录", "页码", "客户填写"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, len(headers))
    for idx, f in enumerate(extraction.fill_items, start=1):
        r = idx + 1
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=f.section)
        ws.cell(row=r, column=3, value=f.field_name)
        ws.cell(row=r, column=4, value=f.description)
        ws.cell(row=r, column=5, value="是" if f.required else "否")
        ws.cell(row=r, column=6, value=f.quote)
        ws.cell(row=r, column=7, value=f.page or "")
        ws.cell(row=r, column=8, value="")
    if not extraction.fill_items:
        ws.cell(row=2, column=1, value=1)
        ws.cell(row=2, column=3, value="报价金额（小写/大写）")
        ws.cell(row=2, column=4, value="按招标文件报价一览表")
        ws.cell(row=2, column=5, value="是")

    widths = [6, 14, 22, 28, 8, 40, 8, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    if meta and meta.review and meta.review.issues:
        wsr = wb.create_sheet("审查记录")
        wsr.cell(row=1, column=1, value="严重程度")
        wsr.cell(row=1, column=2, value="代码")
        wsr.cell(row=1, column=3, value="说明")
        wsr.cell(row=1, column=4, value="字段")
        _style_header(wsr, 1, 4)
        for i, iss in enumerate(meta.review.issues, start=1):
            wsr.cell(row=i + 1, column=1, value=iss.severity)
            wsr.cell(row=i + 1, column=2, value=iss.code)
            wsr.cell(row=i + 1, column=3, value=iss.message)
            wsr.cell(row=i + 1, column=4, value=iss.field_path or "")
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
