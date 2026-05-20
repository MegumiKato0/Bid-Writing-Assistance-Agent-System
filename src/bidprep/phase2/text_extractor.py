from __future__ import annotations

import re
from pathlib import Path

import fitz  # PyMuPDF

from bidprep.pdf_text import extract_pages, pages_to_full_text
from bidprep.phase2.ocr_service import ocr_pdf_page_raster, ocr_service_available
from bidprep.phase2.schemas import ClientFile

# 文本过少时视为扫描件或版式 PDF，触发 OCR 兜底
_MIN_CHARS_PDF = 80
_MIN_CHARS_PER_PAGE_RATIO = 12


def _docx_paragraphs(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf_text_and_pages(path: Path) -> tuple[str, int, bool]:
    pages = extract_pages(str(path))
    full = pages_to_full_text(pages)
    n = len(pages) or 1
    char_count = len(re.sub(r"\s+", "", full))
    thin = char_count < _MIN_CHARS_PDF or (char_count // max(n, 1)) < _MIN_CHARS_PER_PAGE_RATIO
    return full, len(pages), thin


def extract_text_for_client(
    client: ClientFile,
    *,
    ocr_fallback: bool = True,
) -> tuple[str, bool]:
    """
    返回 (全文, ocr_used)。
    优先直接抽取；PDF 文本极薄时对每页尝试 OCR（若 ocr 不可用则返回薄文本）。
    """
    path = Path(client.abs_path)
    try:
        if not path.is_file() or path.stat().st_size == 0:
            client.page_count = 0
            return "", False
    except OSError:
        client.page_count = 0
        return "", False

    ext = client.extension.lower()
    ocr_used = False

    if ext == ".pdf":
        text, page_count, thin = _pdf_text_and_pages(path)
        client.page_count = page_count
        client.is_scanned_pdf = thin
        if thin and ocr_fallback and ocr_service_available():
            ocr_parts: list[str] = []
            for i in range(page_count):
                t = ocr_pdf_page_raster(str(path), i)
                if t.strip():
                    ocr_parts.append(f"【第{i + 1}页 OCR】\n{t.strip()}")
            if ocr_parts:
                text = "\n\n".join(ocr_parts)
                ocr_used = True
        return text, ocr_used

    if ext == ".docx":
        try:
            text = _docx_paragraphs(path)
        except Exception:
            text = ""
        client.page_count = max(1, text.count("\n") // 40 + 1) if text else 1
        return text, False

    if ext in (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"):
        client.page_count = 1
        if ocr_fallback and ocr_service_available():
            from bidprep.phase2.ocr_service import ocr_image_file

            t = ocr_image_file(str(path))
            return t, bool(t.strip())
        return "", False

    return "", False


def hydrate_clients_text(clients: list[ClientFile], *, ocr_fallback: bool = True) -> None:
    for c in clients:
        text, ocr = extract_text_for_client(c, ocr_fallback=ocr_fallback)
        c.extracted_text = text
        c.ocr_used = ocr
