from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bidprep.phase2.file_profile_schemas import Phase2CoverageSnapshot
from bidprep.phase2.schemas import ArchiveIndexEntry, ClientFile, FileUsageResult, RequirementVerification, VerificationStatus

HEADER_FILL = PatternFill(start_color="66CCFF", end_color="66CCFF", fill_type="solid")
HEADER_FONT = Font(bold=True, color="000000")


def _style_header(ws, headers: list[str]) -> None:
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_verification_xlsx(
    rows: list[RequirementVerification],
    path: Path,
    *,
    coverage: Phase2CoverageSnapshot | None = None,
    clients_by_id: dict[str, ClientFile] | None = None,
    file_usage: list[FileUsageResult] | None = None,
    include_unused_files: bool = True,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "核验总表"
    headers = [
        "序号",
        "材料类别",
        "标准材料名称",
        "招标状态",
        "优先级",
        "核验状态",
        "命中文件ID",
        "候选文件ID",
        "替代文件ID",
        "条件满足",
        "归档文件名",
        "PNG路径",
        "关键信息",
        "备注",
    ]
    _style_header(sheet, headers)
    for row_index, row in enumerate(rows, start=2):
        values = [
            row.ordinal,
            row.category,
            row.name,
            row.tender_status,
            row.priority.value,
            row.status.value,
            ", ".join(row.matched_file_ids),
            ", ".join(row.probable_file_ids),
            ", ".join(row.substitution_source_file_ids),
            "是" if row.conditional else "否",
            "；".join(row.archive_file_names),
            "；".join(row.png_paths),
            "\n".join(row.key_information),
            row.notes,
        ]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col_index, width in enumerate([8, 16, 28, 16, 12, 16, 24, 24, 24, 10, 32, 36, 40, 36], start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    if file_usage is not None:
        usage_sheet = workbook.create_sheet("文件使用情况")
        usage_headers = ["文件ID", "原文件名", "分类", "匹配到的 requirement", "使用状态"]
        _style_header(usage_sheet, usage_headers)
        for row_index, item in enumerate(file_usage, start=2):
            values = [
                item.file_id,
                item.original_name,
                item.doc_category or "",
                "；".join(item.matched_requirement_names),
                item.status,
            ]
            for col_index, value in enumerate(values, start=1):
                usage_sheet.cell(row=row_index, column=col_index, value=value)
        for col_index, width in enumerate([22, 36, 18, 48, 16], start=1):
            usage_sheet.column_dimensions[get_column_letter(col_index)].width = width

    if coverage is not None and include_unused_files:
        unused_sheet = workbook.create_sheet("未使用文件")
        _style_header(unused_sheet, ["文件ID", "原文件名"])
        if clients_by_id:
            for row_index, file_id in enumerate(coverage.program_unused_file_ids, start=2):
                unused_sheet.cell(row=row_index, column=1, value=file_id)
                unused_sheet.cell(row=row_index, column=2, value=clients_by_id.get(file_id).original_filename if file_id in clients_by_id else "")

    workbook.save(path)
    workbook.close()


def write_missing_md(
    rows: list[RequirementVerification],
    path: Path,
    *,
    file_usage: list[FileUsageResult] | None = None,
    clients_by_id: dict[str, ClientFile] | None = None,
    include_unused_files: bool = True,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = ["# 阶段二核验缺口", ""]
    for row in rows:
        if row.status in {VerificationStatus.matched, VerificationStatus.substituted}:
            continue
        lines.append(f"- {row.ordinal}. {row.name} [{row.status.value}]")
        if row.notes:
            lines.append(f"  备注: {row.notes}")
        if row.candidate_file_ids:
            lines.append("  候选文件: " + ", ".join(row.candidate_file_ids))
        lines.append("")
    if include_unused_files and file_usage is not None and clients_by_id is not None:
        lines.extend(["## 未使用文件", ""])
        for item in file_usage:
            if item.status != "unused":
                continue
            lines.append(f"- {item.file_id}: {clients_by_id[item.file_id].original_filename}")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_coverage_snapshot_json(snapshot: Phase2CoverageSnapshot, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(snapshot.model_dump(), ensure_ascii=False, indent=2), encoding="utf-8")


def write_unused_and_unmatched_md(snapshot: Phase2CoverageSnapshot, clients_by_id: dict[str, ClientFile], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# 未使用文件与未闭环 requirement",
        "",
        "## 未使用文件",
        "",
    ]
    for file_id in snapshot.program_unused_file_ids:
        client = clients_by_id.get(file_id)
        lines.append(f"- {file_id}: {client.original_filename if client else ''}")
    lines.extend(["", "## 未闭环 requirement", ""])
    for requirement_id in snapshot.unmatched_requirement_ids:
        lines.append(f"- {requirement_id}")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_archive_index_json(entries: dict[str, ArchiveIndexEntry], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({k: v.model_dump() for k, v in entries.items()}, ensure_ascii=False, indent=2), encoding="utf-8")


def write_file_usage_json(file_usage: list[FileUsageResult], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps([item.model_dump() for item in file_usage], ensure_ascii=False, indent=2), encoding="utf-8")


def write_verification_json(rows: list[RequirementVerification], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps([row.model_dump() for row in rows], ensure_ascii=False, indent=2), encoding="utf-8")
