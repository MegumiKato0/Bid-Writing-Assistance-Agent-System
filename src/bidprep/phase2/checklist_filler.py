from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bidprep.phase2.schemas import ArchiveIndexEntry, ClientFile, Requirement, RequirementVerification

HEADER_FILL = PatternFill(start_color="66CCFF", end_color="66CCFF", fill_type="solid")
HEADER_FONT = Font(bold=True, color="000000")

FILL_HEADERS = [
    "标准材料名称",
    "当前状态",
    "匹配文件名",
    "归档文件名",
    "PNG路径",
    "提取到的关键信息",
    "备注",
]


def _style_header(ws, row: int, start_col: int, headers: list[str]) -> None:
    for offset, header in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + offset, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _status_label(verification: RequirementVerification) -> str:
    return verification.status.value


def _collect_key_information(file_ids: list[str], clients_by_id: dict[str, ClientFile]) -> str:
    info: list[str] = []
    for file_id in file_ids:
        client = clients_by_id.get(file_id)
        if client is None or client.document_profile is None:
            continue
        profile = client.document_profile
        lines = list(profile.key_information)
        if profile.llm_summary:
            lines.append(profile.llm_summary)
        if lines:
            info.append(f"{client.original_filename}: " + " | ".join(lines[:6]))
    return "\n".join(info)


def _join_archive_names(file_ids: list[str], archive_index: dict[str, ArchiveIndexEntry]) -> str:
    items: list[str] = []
    for file_id in file_ids:
        entry = archive_index.get(file_id)
        if entry is None:
            continue
        items.extend(entry.standardized_filenames)
    return "；".join(dict.fromkeys(items))


def _join_png_paths(file_ids: list[str], archive_index: dict[str, ArchiveIndexEntry]) -> str:
    items: list[str] = []
    for file_id in file_ids:
        entry = archive_index.get(file_id)
        if entry is None:
            continue
        items.extend(entry.preview_png_paths)
    return "；".join(dict.fromkeys(items))


def _remarks(verification: RequirementVerification) -> str:
    notes: list[str] = []
    if verification.notes:
        notes.append(verification.notes)
    if verification.substitution_applied:
        notes.append("已走替代满足逻辑")
    if verification.conditional:
        notes.append(verification.condition_note or "条件满足/条件核验")
    return "\n".join(notes)


def _write_row(
    ws,
    row_index: int,
    start_col: int,
    requirement: Requirement,
    verification: RequirementVerification,
    clients_by_id: dict[str, ClientFile],
    archive_index: dict[str, ArchiveIndexEntry],
) -> None:
    file_ids = list(
        dict.fromkeys(
            verification.matched_file_ids + verification.probable_file_ids + verification.substitution_source_file_ids
        )
    )
    matched_names = "；".join(
        client.original_filename for file_id in file_ids if (client := clients_by_id.get(file_id)) is not None
    )
    values = [
        requirement.name,
        _status_label(verification),
        matched_names,
        _join_archive_names(file_ids, archive_index),
        _join_png_paths(file_ids, archive_index),
        _collect_key_information(file_ids, clients_by_id),
        _remarks(verification),
    ]
    for offset, value in enumerate(values):
        cell = ws.cell(row=row_index, column=start_col + offset, value=value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_filled_materials_checklist(
    *,
    baseline_xlsx: Path | None,
    requirements: list[Requirement],
    verifications: list[RequirementVerification],
    clients_by_id: dict[str, ClientFile],
    archive_index: dict[str, ArchiveIndexEntry],
    out_path: Path,
) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    verification_by_id = {verification.requirement_id: verification for verification in verifications}
    requirement_by_ordinal = {requirement.ordinal: requirement for requirement in requirements}
    requirement_by_name = {requirement.name.strip(): requirement for requirement in requirements}

    if baseline_xlsx and baseline_xlsx.is_file() and baseline_xlsx.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(baseline_xlsx)
        worksheet = workbook["材料清单"] if "材料清单" in workbook.sheetnames else workbook.active
        start_col = worksheet.max_column + 1
        _style_header(worksheet, 1, start_col, FILL_HEADERS)

        for row_index in range(2, worksheet.max_row + 1):
            ordinal_raw = worksheet.cell(row=row_index, column=1).value
            name_raw = worksheet.cell(row=row_index, column=3).value
            try:
                ordinal = int(ordinal_raw) if ordinal_raw not in (None, "") else 0
            except (TypeError, ValueError):
                ordinal = 0
            name = str(name_raw).strip() if name_raw else ""
            if not name:
                continue
            requirement = requirement_by_ordinal.get(ordinal) or requirement_by_name.get(name)
            if requirement is None:
                continue
            verification = verification_by_id.get(requirement.id)
            if verification is None:
                continue
            _write_row(worksheet, row_index, start_col, requirement, verification, clients_by_id, archive_index)

        for col in range(start_col, start_col + len(FILL_HEADERS)):
            worksheet.column_dimensions[get_column_letter(col)].width = 28
        workbook.save(out_path)
        workbook.close()
        return

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "材料清单_回填"
    base_headers = ["序号", "材料类别", "材料名称", "招标状态", "备注"]
    _style_header(worksheet, 1, 1, base_headers + FILL_HEADERS)
    for row_index, requirement in enumerate(requirements, start=2):
        verification = verification_by_id.get(requirement.id)
        worksheet.cell(row=row_index, column=1, value=requirement.ordinal)
        worksheet.cell(row=row_index, column=2, value=requirement.category)
        worksheet.cell(row=row_index, column=3, value=requirement.name)
        worksheet.cell(row=row_index, column=4, value=requirement.tender_status)
        worksheet.cell(row=row_index, column=5, value=requirement.remark)
        if verification:
            _write_row(worksheet, row_index, 6, requirement, verification, clients_by_id, archive_index)

    for column in range(1, 5 + len(FILL_HEADERS) + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 22 if column != 3 else 36
    workbook.save(out_path)
    workbook.close()
