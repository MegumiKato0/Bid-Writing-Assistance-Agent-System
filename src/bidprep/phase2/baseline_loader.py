from __future__ import annotations

import json
import re
from enum import IntEnum
from pathlib import Path
from openpyxl import load_workbook

from bidprep.phase2.schemas import (
    BaselineFormat,
    BaselineStandardDocument,
    Requirement,
    RequirementPriority,
)


# ---------------------------------------------------------------------------
# 多文件策略（与产品说明一致）
# ---------------------------------------------------------------------------
# 优先级（从高到低）：JSON > xlsx/xlsm > PDF > Markdown
# 系统在最高 nonempty 档中必须恰好 1 个候选；同档多个则报错提示用户删减。
# extraction_meta.json 不参与 JSON 档（无 materials/requirements），视为跳过。
# ---------------------------------------------------------------------------
# xlsx 解析规则
# ---------------------------------------------------------------------------
# - 优先工作表「材料清单」，否则 active
# - 表头须含「材料名称」列；推荐列：序号、材料类别、材料名称、是否具备、备注
# - 行：跳过空行；材料名称为空或「—」则跳过
# ---------------------------------------------------------------------------


class _ClassifyResult(IntEnum):
    skip = 0
    json = 1
    xlsx = 2
    pdf = 3
    md = 4
    unsupported = 99


def _classify_path(path: Path) -> _ClassifyResult:
    suf = path.suffix.lower()
    name = path.name.lower()
    if suf == ".json":
        if name == "extraction_meta.json":
            return _ClassifyResult.skip
        return _ClassifyResult.json
    if suf in (".xlsx", ".xlsm"):
        return _ClassifyResult.xlsx
    if suf == ".pdf":
        return _ClassifyResult.pdf
    if suf in (".md", ".markdown"):
        return _ClassifyResult.md
    return _ClassifyResult.unsupported


def _format_label(c: _ClassifyResult) -> str:
    return {
        _ClassifyResult.json: "JSON",
        _ClassifyResult.xlsx: "Excel（.xlsx/.xlsm）",
        _ClassifyResult.pdf: "PDF",
        _ClassifyResult.md: "Markdown（.md）",
    }.get(c, "未知")


def select_baseline_file(paths: list[Path]) -> tuple[Path, BaselineFormat, str]:
    """
    在多个上传路径中按优先级选一个基准文件。

    返回 (路径, 格式枚举, 人类可读说明)。
    """
    if not paths:
        raise ValueError("至少需要一个基准清单文件")

    bad = [p for p in paths if _classify_path(p) == _ClassifyResult.unsupported]
    if bad:
        names = "、".join(p.name for p in bad)
        raise ValueError(
            f"不支持的基准扩展名（仅 .json / .xlsx / .xlsm / .pdf / .md）：{names}"
        )

    skipped_meta = [p.name for p in paths if _classify_path(p) == _ClassifyResult.skip]
    candidates = [p for p in paths if _classify_path(p) != _ClassifyResult.skip]
    if not candidates:
        raise ValueError(
            "未找到可用基准清单：已忽略 extraction_meta.json（不含 materials/requirements）。"
            "请上传投标材料清单表 .xlsx / .pdf / .md，或含清单的 baseline .json。"
        )

    tier_order = (
        _ClassifyResult.json,
        _ClassifyResult.xlsx,
        _ClassifyResult.pdf,
        _ClassifyResult.md,
    )
    for tier in tier_order:
        group = [p for p in candidates if _classify_path(p) == tier]
        if not group:
            continue
        if len(group) > 1:
            names = "、".join(p.name for p in group)
            raise ValueError(
                f"同优先级存在多个基准文件（{_format_label(tier)}）：{names}。"
                f"规则为 JSON > Excel > PDF > Markdown，且每一档只能选 1 个文件；请删除或取消多余的同格式文件后再试。"
            )
        chosen = group[0]
        bf = {
            _ClassifyResult.json: BaselineFormat.json,
            _ClassifyResult.xlsx: BaselineFormat.xlsx,
            _ClassifyResult.pdf: BaselineFormat.pdf,
            _ClassifyResult.md: BaselineFormat.md,
        }[tier]
        note = f"基准清单：{chosen.name}（{_format_label(tier)}，优先级 JSON > Excel > PDF > Markdown）"
        if skipped_meta:
            note += f"；已忽略 {', '.join(skipped_meta)}"
        return chosen, bf, note

    raise ValueError("内部错误：未能为基准文件分级")


def _priority_from_tender_status(status: str) -> RequirementPriority:
    s = (status or "").strip()
    if not s:
        return RequirementPriority.required
    if "不需要" in s or "无须" in s or "无需" in s or "不适用" in s:
        return RequirementPriority.exempt
    if "可选" in s or "如有" in s or "若有" in s:
        return RequirementPriority.optional
    return RequirementPriority.required


def raw_dicts_to_requirements(raw_items: list) -> list[Requirement]:
    """将 JSON 数组项（dict）转为 Requirement。"""
    out: list[Requirement] = []
    for i, item in enumerate(raw_items, start=1):
        if not isinstance(item, dict):
            continue
        ordinal = int(item.get("ordinal") or item.get("序号") or i)
        name = str(item.get("name") or item.get("材料名称") or "").strip()
        if not name:
            continue
        category = str(item.get("category") or item.get("材料类别") or "")
        tender_status = str(
            item.get("tender_status") or item.get("是否具备") or item.get("是否已准备") or ""
        )
        remark = str(item.get("remark") or item.get("备注") or "")
        rid = str(item.get("id") or f"R{ordinal:03d}")
        out.append(
            Requirement(
                id=rid,
                ordinal=ordinal,
                category=category,
                name=name,
                tender_status=tender_status,
                remark=remark,
                priority=_priority_from_tender_status(tender_status),
            )
        )
    return out


def normalize_payload_to_requirements(data: object) -> list[Requirement]:
    """
    将任意支持的 JSON 根结构转为 Requirement 列表，并等价于标准包 ``{"requirements": [...]}`` 的条目部分。
    """
    if isinstance(data, list):
        return raw_dicts_to_requirements(data)
    if not isinstance(data, dict):
        raise ValueError("JSON 根须为对象或数组")
    materials = data.get("materials") or data.get("requirements")
    if not isinstance(materials, list):
        raise ValueError(
            "JSON 需包含 materials 或 requirements 数组（extraction_meta.json 不是清单基准，请换用清单表）"
        )
    return raw_dicts_to_requirements(materials)


def requirements_to_standard_document(requirements: list[Requirement]) -> BaselineStandardDocument:
    """统一标准结构（供序列化/调试）。"""
    rows = [r.model_dump(mode="json") for r in requirements]
    return BaselineStandardDocument(requirements=rows)


def load_baseline_from_json(path: str | Path) -> list[Requirement]:
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    return normalize_payload_to_requirements(raw)


def load_baseline_from_xlsx(path: str | Path) -> list[Requirement]:
    """从 Phase 1 导出的「投标材料清单表.xlsx」工作表「材料清单」加载基准行。"""
    p = Path(path)
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        if "材料清单" in wb.sheetnames:
            ws = wb["材料清单"]
        else:
            ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]

    col_idx = {h: i for i, h in enumerate(header)}

    def _col(name: str, alt: tuple[str, ...] = ()) -> int | None:
        for k in (name,) + alt:
            if k in col_idx:
                return col_idx[k]
        return None

    i_ord = _col("序号")
    i_cat = _col("材料类别")
    i_name = _col("材料名称")
    i_stat = _col("是否具备", ("是否已准备",))
    i_remark = _col("备注")
    if i_name is None:
        raise ValueError("基准表缺少「材料名称」列，请使用投标材料清单表.xlsx")

    out: list[Requirement] = []
    seen_ordinals: dict[int, int] = {}
    for row in rows[1:]:
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue

        def cell(idx: int | None) -> str:
            if idx is None or idx >= len(row):
                return ""
            v = row[idx]
            if v is None:
                return ""
            return str(v).strip()

        ord_raw = cell(i_ord)
        name = cell(i_name)
        if not name or name in ("—", "-", "..."):
            continue
        ordinal = 0
        if ord_raw:
            m = re.search(r"\d+", ord_raw)
            if m:
                ordinal = int(m.group())
        if ordinal <= 0:
            ordinal = len(out) + 1
        tender_status = cell(i_stat)
        remark = cell(i_remark)
        category = cell(i_cat)
        rid = f"R{ordinal:03d}"
        if rid in {r.id for r in out}:
            seen_ordinals[ordinal] = seen_ordinals.get(ordinal, 0) + 1
            rid = f"R{ordinal:03d}_{seen_ordinals[ordinal]}"
        out.append(
            Requirement(
                id=rid,
                ordinal=ordinal,
                category=category,
                name=name,
                tender_status=tender_status,
                remark=remark,
                priority=_priority_from_tender_status(tender_status),
            )
        )
    return out


def _split_table_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    parts = [c.strip() for c in s.split("|")]
    return parts


def _is_md_separator_row(line: str) -> bool:
    t = line.replace(" ", "")
    return bool(re.match(r"^\|?[\-:|]+\|?$", t))


def load_baseline_from_md(path: str | Path) -> list[Requirement]:
    """从「投标材料准备清单」类 Markdown 中的表格解析（含 序号/材料名称/是否已准备/备注 等列）。"""
    text = Path(path).read_text(encoding="utf-8")
    lines = text.splitlines()
    header_idx = -1
    for i, ln in enumerate(lines):
        if "|" not in ln:
            continue
        low = ln.replace(" ", "")
        if "材料名称" in low and ("序号" in low or "序號" in low):
            header_idx = i
            break
    if header_idx < 0:
        raise ValueError("Markdown 中未找到含「序号」「材料名称」的表格表头")

    header_cells = _split_table_row(lines[header_idx])
    norm = [re.sub(r"\s+", "", c) for c in header_cells]

    def find_col(*cands: str) -> int | None:
        for j, h in enumerate(norm):
            for c in cands:
                if c in h:
                    return j
        return None

    i_ord = find_col("序号", "序號")
    i_name = find_col("材料名称", "材料名稱")
    if i_name is None:
        raise ValueError("Markdown 表缺少「材料名称」列")
    i_stat = find_col("是否具备", "是否已准备", "是否準備")
    i_rem = find_col("备注", "備註")
    i_cat = find_col("材料类别", "材料類別", "类别", "類別")

    out: list[Requirement] = []
    seen_ordinals: dict[int, int] = {}
    for ln in lines[header_idx + 1 :]:
        s = ln.strip()
        if not s or "|" not in s:
            continue
        if _is_md_separator_row(s):
            continue
        cells = _split_table_row(s)
        if len(cells) < 2:
            continue

        def gc(idx: int | None) -> str:
            if idx is None or idx >= len(cells):
                return ""
            return cells[idx].strip()

        name = gc(i_name)
        if not name or name in ("—", "-", "..."):
            continue
        ord_raw = gc(i_ord)
        ordinal = 0
        if ord_raw:
            m = re.search(r"\d+", ord_raw)
            if m:
                ordinal = int(m.group())
        if ordinal <= 0:
            ordinal = len(out) + 1
        tender_status = gc(i_stat)
        remark = gc(i_rem)
        category = gc(i_cat)
        rid = f"R{ordinal:03d}"
        if rid in {r.id for r in out}:
            seen_ordinals[ordinal] = seen_ordinals.get(ordinal, 0) + 1
            rid = f"R{ordinal:03d}_{seen_ordinals[ordinal]}"
        out.append(
            Requirement(
                id=rid,
                ordinal=ordinal,
                category=category,
                name=name,
                tender_status=tender_status,
                remark=remark,
                priority=_priority_from_tender_status(tender_status),
            )
        )
    return out


def load_baseline_from_pdf(path: str | Path) -> list[Requirement]:
    """
    从「投标材料清单表」PDF（与阶段一导出同源）抽取文本并解析为行。
    依赖表头行含「材料名称」；数据行以序号数字开头。
    """
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ValueError("解析 PDF 基准需要 PyMuPDF（pymupdf）") from e

    p = Path(path)
    doc = fitz.open(p)
    try:
        text = "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()

    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    hi = -1
    for i, ln in enumerate(lines):
        if "材料名称" in ln.replace(" ", "") and ("序号" in ln.replace(" ", "") or "序號" in ln.replace(" ", "")):
            hi = i
            break
    if hi < 0:
        raise ValueError("PDF 中未识别到清单表头（需含「序号」「材料名称」列，建议使用阶段一导出的投标材料清单表.pdf）")

    out: list[Requirement] = []
    seen_ordinals: dict[int, int] = {}
    for ln in lines[hi + 1 :]:
        if "投标材料清单表" in ln and len(ln) < 24:
            continue
        m = re.match(r"^(\d+)\s+(.+)$", ln)
        if not m:
            continue
        ordinal = int(m.group(1))
        rest = m.group(2).strip()
        parts = re.split(r"\s{2,}", rest)
        category = ""
        if len(parts) >= 4:
            category, name, tender_status, remark = (
                parts[0],
                parts[1],
                parts[2],
                " ".join(parts[3:]).strip(),
            )
        elif len(parts) == 3:
            name, tender_status, remark = parts[0], parts[1], parts[2]
        elif len(parts) == 2:
            name, tender_status, remark = parts[0], parts[1], ""
        else:
            name, tender_status, remark = rest, "", ""

        if not name or name in ("—", "-", "..."):
            continue
        rid = f"R{ordinal:03d}"
        if rid in {r.id for r in out}:
            seen_ordinals[ordinal] = seen_ordinals.get(ordinal, 0) + 1
            rid = f"R{ordinal:03d}_{seen_ordinals[ordinal]}"
        out.append(
            Requirement(
                id=rid,
                ordinal=ordinal,
                category=category,
                name=name,
                tender_status=tender_status,
                remark=remark,
                priority=_priority_from_tender_status(tender_status),
            )
        )
    if not out:
        raise ValueError("PDF 中未解析出任何材料行，请确认是否为阶段一导出的「投标材料清单表.pdf」")
    return out


def load_baseline(path: str | Path) -> list[Requirement]:
    """按扩展名解析单个基准文件 → Requirement 列表（内部已等价于标准 requirements 结构）。"""
    p = Path(path)
    suf = p.suffix.lower()
    if suf == ".json":
        return load_baseline_from_json(p)
    if suf in (".xlsx", ".xlsm"):
        return load_baseline_from_xlsx(p)
    if suf == ".pdf":
        return load_baseline_from_pdf(p)
    if suf in (".md", ".markdown"):
        return load_baseline_from_md(p)
    raise ValueError(f"不支持的基准文件类型：{suf}")


def resolve_and_load_baselines(paths: list[Path]) -> tuple[list[Requirement], str, Path]:
    """
    多路径时按优先级选唯一基准文件并加载。

    返回 (requirements, 日志说明, 被选中的基准文件路径) — 供回填清单表时复制原表版式。
    """
    chosen, _fmt, note = select_baseline_file(paths)
    try:
        reqs = load_baseline(chosen)
    except ValueError as e:
        raise ValueError(f"{chosen.name}：{e}") from e
    return reqs, note, chosen


