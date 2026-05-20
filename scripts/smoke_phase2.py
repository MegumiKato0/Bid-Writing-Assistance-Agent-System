# -*- coding: utf-8 -*-
from __future__ import annotations

import shutil
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
_SRC = _ROOT / "src"
if _SRC.is_dir() and str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from bidprep.llm_runtime import runtime_from_settings  # noqa: E402
from bidprep.phase2.pipeline import run_phase2_verify  # noqa: E402


def main() -> None:
    tmp = _ROOT / "_phase2_smoke"
    tmp.mkdir(exist_ok=True)
    base = tmp / "baseline.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "材料清单"
    ws.append(["序号", "材料类别", "材料名称", "是否具备", "备注"])
    ws.append([1, "资格", "营业执照等证明文件", "☐ 需准备", ""])
    ws.append([2, "资格", "依法缴纳税收证明材料", "☐ 需准备", ""])
    ws.append([3, "资格", "资格承诺函（信用承诺制）", "☐ 需准备", ""])
    wb.save(base)

    docx = tmp / "营业执照示例.docx"
    d = Document()
    d.add_paragraph("营业执照 统一社会信用代码 有限公司")
    d.add_paragraph("供应商资格承诺函（信用承诺制）")
    d.save(str(docx))

    out = tmp / "out"
    if out.exists():
        shutil.rmtree(out)
    run_phase2_verify(
        base,
        docx,
        out,
        ocr_fallback=False,
        progress=print,
        llm=runtime_from_settings(),
    )
    print("reports:", list((out / "reports").glob("*")))


if __name__ == "__main__":
    main()
